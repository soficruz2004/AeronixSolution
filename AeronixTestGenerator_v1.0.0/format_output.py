try:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    print("Warning: python-docx not installed. Install with: pip install python-docx")

import os
from datetime import datetime
from typing import Dict, List, Any

class ProfessionalFormatter:
    def __init__(self):
        if not DOCX_AVAILABLE:
            raise ImportError("python-docx library not available. Install with: pip install python-docx")
        self.doc = Document()
    
    def create_test_document(self, test_data: str, device_type: str = "PCB") -> str:
        """Create professional test document"""
        # Header
        self.add_header(device_type)
        
        # Parse and format test data
        test_steps = self.parse_test_data(test_data)
        
        # Add overview
        self.add_overview(device_type, len(test_steps))
        
        # Add test steps
        self.add_test_steps(test_steps)
        
        # Add footer
        self.add_footer()
        
        # Save document
        filename = f"output/{device_type}_Test_Procedure_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        os.makedirs("output", exist_ok=True)
        self.doc.save(filename)
        return filename
    
    def add_header(self, device_type: str):
        """Add document header"""
        title = self.doc.add_paragraph(f"{device_type} Test Procedure")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in title.runs:
            run.font.size = Pt(18)
            run.font.bold = True
        
        table = self.doc.add_table(rows=4, cols=2)
        table.style = 'Table Grid'
        
        cells = [
            ("Document ID:", f"AE{datetime.now().strftime('%y%m%d')}-001"),
            ("Generated:", datetime.now().strftime("%Y-%m-%d %H:%M")),
            ("Device Type:", device_type),
            ("Status:", "DRAFT")
        ]
        
        for i, (label, value) in enumerate(cells):
            table.cell(i, 0).text = label
            table.cell(i, 1).text = value
        
        self.doc.add_paragraph()
    
    def add_overview(self, device_type: str, step_count: int):
        """Add test overview section"""
        heading = self.doc.add_paragraph("Overview")
        for run in heading.runs:
            run.font.size = Pt(14)
            run.font.bold = True
        
        overview_text = f"""This document provides comprehensive test procedures for the {device_type} device. 
The testing covers power validation, functional verification, and performance characterization.

Test Coverage:
• Power supply validation
• Digital interface testing  
• RF performance (if applicable)
• Environmental compliance
• Safety verification

Total Test Steps: {step_count}
Estimated Duration: {step_count * 15} minutes
"""
        self.doc.add_paragraph(overview_text)
        self.doc.add_paragraph()
    
    def parse_test_data(self, test_data: str) -> List[Dict]:
        """Parse AI-generated test data into structured format"""
        steps = []
        current_step = {}
        
        for line in test_data.split('\n'):
            line = line.strip()
            if not line:
                continue
                
            if line.startswith(('1. Title:', '2. Title:', '3. Title:')):
                if current_step:
                    steps.append(current_step)
                current_step = {'title': line.split(':', 1)[1].strip()}
            elif line.startswith(('2. Type:', '3. Type:')):
                current_step['type'] = line.split(':', 1)[1].strip()
            elif line.startswith(('3. Description:', '4. Description:')):
                current_step['description'] = line.split(':', 1)[1].strip()
            elif line.startswith(('4. Expected:', '5. Expected:')):
                current_step['expected'] = line.split(':', 1)[1].strip()
            elif line.startswith(('5. Equipment:', '6. Equipment:')):
                current_step['equipment'] = line.split(':', 1)[1].strip()
            elif line.startswith(('6. Points:', '7. Points:')):
                current_step['points'] = line.split(':', 1)[1].strip()
        
        if current_step:
            steps.append(current_step)
        
        return steps
    
    def add_test_steps(self, steps: List[Dict]):
        """Add formatted test steps"""
        heading = self.doc.add_paragraph("Test Procedures")
        for run in heading.runs:
            run.font.size = Pt(14)
            run.font.bold = True
        
        for i, step in enumerate(steps, 1):
            # Step header
            step_title = f"Step {i}: {step.get('title', 'Untitled Test')}"
            step_para = self.doc.add_paragraph(step_title)
            for run in step_para.runs:
                run.font.bold = True
            
            # Step details table
            table = self.doc.add_table(rows=5, cols=2)
            table.style = 'Table Grid'
            
            details = [
                ("Type:", step.get('type', 'N/A')),
                ("Description:", step.get('description', 'N/A')),
                ("Expected Result:", step.get('expected', 'N/A')),
                ("Equipment:", step.get('equipment', 'N/A')),
                ("Test Points:", step.get('points', 'N/A'))
            ]
            
            for j, (label, value) in enumerate(details):
                table.cell(j, 0).text = label
                table.cell(j, 1).text = value
            
            # Results section
            results_table = self.doc.add_table(rows=2, cols=3)
            results_table.style = 'Table Grid'
            
            headers = ["Result", "Notes", "Signature"]
            for k, header in enumerate(headers):
                results_table.cell(0, k).text = header
            
            results_table.cell(1, 0).text = "☐ PASS  ☐ FAIL"
            
            self.doc.add_paragraph()
    
    def add_footer(self):
        """Add document footer"""
        self.doc.add_paragraph()
        footer_text = f"""
Generated by Aeronix Test Generator
© {datetime.now().year} Vibecoders. Inc
Document Version: 1.0
"""
        footer = self.doc.add_paragraph(footer_text)
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

def format_test_output(test_data: str, device_type: str = "PCB") -> str:
    """Main function to format test output"""
    if not DOCX_AVAILABLE:
        return create_text_output(test_data, device_type)
    try:
        formatter = ProfessionalFormatter()
        return formatter.create_test_document(test_data, device_type)
    except Exception as e:
        print(f"Error creating Word document: {e}")
        return create_text_output(test_data, device_type)

def create_text_output(test_data: str, device_type: str = "PCB") -> str:
    """Fallback text output when docx is not available"""
    filename = f"output/{device_type}_Test_Procedure_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"
    os.makedirs("output", exist_ok=True)
    
    with open(filename, 'w') as f:
        f.write(f"{device_type} Test Procedure\n")
        f.write("=" * 50 + "\n\n")
        f.write(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}\n")
        f.write(f"Device: {device_type}\n\n")
        f.write("Test Data:\n")
        f.write("-" * 20 + "\n")
        f.write(test_data)
        f.write("\n\n" + "=" * 50)
        f.write("\nGenerated by Aeronix Test Generator")
    
    return filename