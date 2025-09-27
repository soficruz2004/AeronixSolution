from typing import Dict, List, Any, Optional
from dataclasses import dataclass
from pydantic import BaseModel
import csv
import re
import json
import os
import xml.etree.ElementTree as ET

# Optional imports for enhanced file support
try:
    import PyPDF2
    PDF_SUPPORT = True
except ImportError:
    PDF_SUPPORT = False

try:
    from docx import Document
    DOCX_SUPPORT = True
except ImportError:
    DOCX_SUPPORT = False

try:
    import openpyxl
    EXCEL_SUPPORT = True
except ImportError:
    EXCEL_SUPPORT = False

@dataclass
class Component:
    designator: str
    part_number: str
    footprint: str
    quantity: int
    description: str
    test_priority: str = "MEDIUM"
    value: str = ""
    package: str = ""
    supplier: str = ""
    manufacturer: str = ""

@dataclass
class TestPoint:
    net_name: str
    component_ref: str
    x_coord: float
    y_coord: float
    rotation: float = 0.0
    layer: str = "TOP"
    drill_size: float = 0.0

@dataclass
class NetConnection:
    net_name: str
    component_pins: List[str]
    test_points: List[str]

class PCBDataModel(BaseModel):
    bom_components: List[Dict]
    test_points: List[Dict]
    schematics: Dict
    requirements: Dict
    netlist: Dict

class BOMParser:
    def parse(self, content: str, file_path: str = "") -> List[Component]:
        """Parse BOM content based on format detection"""
        file_ext = os.path.splitext(file_path)[1].lower() if file_path else ""
        
        if file_ext == '.bomdoc':
            return self._parse_altium_bom(content, file_path)
        elif content.strip().startswith('<?xml'):
            return self._parse_xml_bom(content)
        elif '\t' in content or ',' in content:
            return self._parse_csv_bom(content)
        elif file_ext in ['.xlsx', '.xls'] and EXCEL_SUPPORT:
            return self._parse_excel_bom(file_path)
        else:
            return self._parse_text_bom(content)

    def _parse_altium_bom(self, content: str, file_path: str) -> List[Component]:
        """Parse Altium BOM document (placeholder - needs actual Altium parser)"""
        components = []
        # Altium .BomDoc files are binary and require specialized parsing
        # For now, return a placeholder component indicating the file was detected
        components.append(Component(
            designator="ALTIUM_BOM",
            part_number="DETECTED",
            footprint="",
            quantity=1,
            description=f"Altium BOM file detected: {os.path.basename(file_path)}",
            test_priority="HIGH"
        ))
        return components

    def _parse_excel_bom(self, file_path: str) -> List[Component]:
        """Parse Excel BOM file"""
        components = []
        try:
            workbook = openpyxl.load_workbook(file_path)
            sheet = workbook.active
            
            # Assume first row is headers
            headers = [cell.value for cell in sheet[1]]
            
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if row[0]:  # Skip empty rows
                    comp = Component(
                        designator=str(row[0]) if row[0] else "",
                        part_number=str(row[1]) if len(row) > 1 and row[1] else "",
                        footprint=str(row[2]) if len(row) > 2 and row[2] else "",
                        quantity=int(row[3]) if len(row) > 3 and isinstance(row[3], (int, float)) else 1,
                        description=str(row[4]) if len(row) > 4 and row[4] else "",
                        value=str(row[5]) if len(row) > 5 and row[5] else "",
                        test_priority=self._assess_component_priority(str(row[0]) if row[0] else "", 
                                                                    str(row[1]) if len(row) > 1 and row[1] else "")
                    )
                    components.append(comp)
        except Exception as e:
            print(f"Error parsing Excel BOM: {e}")
        
        return components

    def _parse_text_bom(self, content: str) -> List[Component]:
        """Parse text-based BOM"""
        components = []
        lines = content.strip().split('\n')
        for line in lines:
            if line.strip().startswith('#') or line.strip() == '':
                continue
            parts = line.split()
            if len(parts) >= 2:
                comp = Component(
                    designator=parts[0].strip(),
                    part_number=parts[1].strip() if len(parts) > 1 else "",
                    footprint=parts[2].strip() if len(parts) > 2 else "",
                    quantity=int(parts[3]) if len(parts) > 3 and parts[3].isdigit() else 1,
                    description=' '.join(parts[4:]) if len(parts) > 4 else "",
                    test_priority=self._assess_component_priority(parts[0], parts[1] if len(parts) > 1 else "")
                )
                components.append(comp)
        return components

    def _parse_xml_bom(self, content: str) -> List[Component]:
        """Parse XML BOM format"""
        components = []
        try:
            root = ET.fromstring(content)
            
            # Look for common XML BOM structures
            for item in root.findall('.//component') or root.findall('.//item') or root.findall('.//part'):
                designator = item.get('designator') or item.get('ref') or item.find('designator')
                if designator is not None:
                    designator = designator.text if hasattr(designator, 'text') else str(designator)
                    
                    part_number = item.get('partnumber') or item.find('partnumber')
                    part_number = part_number.text if hasattr(part_number, 'text') else str(part_number) if part_number else ""
                    
                    comp = Component(
                        designator=designator,
                        part_number=part_number,
                        footprint=self._get_xml_value(item, ['footprint', 'package']),
                        quantity=int(self._get_xml_value(item, ['quantity', 'qty']) or 1),
                        description=self._get_xml_value(item, ['description', 'desc']),
                        value=self._get_xml_value(item, ['value', 'val']),
                        test_priority=self._assess_component_priority(designator, part_number)
                    )
                    components.append(comp)
        except ET.ParseError as e:
            print(f"XML parsing error: {e}")
        
        return components

    def _get_xml_value(self, element, tag_names: List[str]) -> str:
        """Get value from XML element by trying multiple tag names"""
        for tag in tag_names:
            value = element.get(tag) or element.find(tag)
            if value is not None:
                return value.text if hasattr(value, 'text') else str(value)
        return ""

    def _parse_csv_bom(self, content: str) -> List[Component]:
        """Parse CSV BOM format"""
        components = []
        lines = content.strip().split('\n')
        delimiter = '\t' if '\t' in content else ','
        
        reader = csv.reader(lines, delimiter=delimiter)
        headers = next(reader, [])  # Skip header row
        
        for row in reader:
            if len(row) >= 2 and row[0].strip():
                comp = Component(
                    designator=row[0].strip(),
                    part_number=row[1].strip() if len(row) > 1 else "",
                    footprint=row[2].strip() if len(row) > 2 else "",
                    quantity=int(row[3]) if len(row) > 3 and row[3].isdigit() else 1,
                    description=row[4].strip() if len(row) > 4 else "",
                    value=row[5].strip() if len(row) > 5 else "",
                    test_priority=self._assess_component_priority(row[0], row[1] if len(row) > 1 else "")
                )
                components.append(comp)
        return components

    def _assess_component_priority(self, designator: str, part_number: str) -> str:
        """Assess component test priority based on designator and part number"""
        critical_patterns = ['POWER', 'RF', 'CRYSTAL', 'CPU', 'MCU', 'LORA', 'GPS', 'IMU', 'REGULATOR']
        high_priority_refs = ['U', 'IC', 'Q', 'T', 'X', 'Y']
        
        if any(pattern in part_number.upper() for pattern in critical_patterns):
            return "HIGH"
        elif any(pattern in designator.upper() for pattern in critical_patterns):
            return "HIGH"
        elif designator and designator[0] in high_priority_refs:
            return "HIGH"
        elif designator and designator[0] in ['R', 'C', 'L']:
            return "MEDIUM"
        return "LOW"

class CoordinateParser:
    def parse(self, content: str, file_path: str = "") -> List[TestPoint]:
        """Parse coordinate/test point data"""
        file_ext = os.path.splitext(file_path)[1].lower() if file_path else ""
        
        if 'testpoint' in file_path.lower() or 'assembly' in file_path.lower():
            return self._parse_assembly_testpoint_report(content)
        elif content.strip().startswith('D0') or file_ext == '.ipc':
            return self._parse_ipc_format(content)
        elif 'X=' in content and 'Y=' in content:
            return self._parse_xy_format(content)
        else:
            return self._parse_generic_format(content)

    def _parse_assembly_testpoint_report(self, content: str) -> List[TestPoint]:
        """Parse assembly test point reports"""
        test_points = []
        lines = content.strip().split('\n')
        
        for line in lines:
            if line.strip() and not line.startswith('#'):
                # Look for patterns like: TP1 VCC 10.5 20.3 0
                parts = line.split()
                if len(parts) >= 4:
                    try:
                        tp = TestPoint(
                            net_name=parts[1] if len(parts) > 1 else parts[0],
                            component_ref=parts[0],
                            x_coord=float(parts[2]),
                            y_coord=float(parts[3]),
                            rotation=float(parts[4]) if len(parts) > 4 else 0.0
                        )
                        test_points.append(tp)
                    except ValueError:
                        continue
        
        return test_points

    def _parse_ipc_format(self, content: str) -> List[TestPoint]:
        """Parse IPC-356 format"""
        test_points = []
        for line in content.strip().split('\n'):
            if line.startswith('317') or line.startswith('327'):
                parts = line.split()
                if len(parts) >= 3:
                    coords = self._extract_coordinates(parts[2])
                    if coords:
                        tp = TestPoint(
                            net_name=parts[0][3:],  # Remove 317/327
                            component_ref=parts[1],
                            x_coord=coords[0],
                            y_coord=coords[1],
                            rotation=coords[2]
                        )
                        test_points.append(tp)
        return test_points

    def _parse_xy_format(self, content: str) -> List[TestPoint]:
        """Parse X=Y= coordinate format"""
        test_points = []
        for line in content.strip().split('\n'):
            if 'X=' in line and 'Y=' in line:
                coords = self._extract_xy_coordinates(line)
                if coords:
                    tp = TestPoint(
                        net_name=coords.get('net', 'UNKNOWN'),
                        component_ref=coords.get('ref', 'TP'),
                        x_coord=coords['x'],
                        y_coord=coords['y'],
                        rotation=coords.get('rotation', 0.0)
                    )
                    test_points.append(tp)
        return test_points

    def _parse_generic_format(self, content: str) -> List[TestPoint]:
        """Parse generic coordinate format"""
        test_points = []
        for line in content.strip().split('\n'):
            if line.strip() and not line.startswith('#'):
                parts = line.split()
                if len(parts) >= 3:
                    try:
                        tp = TestPoint(
                            net_name=parts[0],
                            component_ref=parts[1] if len(parts) > 1 else parts[0],
                            x_coord=float(parts[-2]),
                            y_coord=float(parts[-1]),
                            rotation=0.0
                        )
                        test_points.append(tp)
                    except (ValueError, IndexError):
                        continue
        return test_points

    def _extract_coordinates(self, coord_string: str) -> Optional[tuple]:
        """Extract coordinates from IPC format string"""
        x_match = re.search(r'X([+-]?\d+)', coord_string)
        y_match = re.search(r'Y([+-]?\d+)', coord_string)
        r_match = re.search(r'R(\d+)', coord_string)
        
        if x_match and y_match:
            x = float(x_match.group(1)) / 1000.0
            y = float(y_match.group(1)) / 1000.0
            rotation = float(r_match.group(1)) if r_match else 0.0
            return (x, y, rotation)
        return None

    def _extract_xy_coordinates(self, line: str) -> Optional[Dict]:
        """Extract coordinates from X=Y= format"""
        x_match = re.search(r'X=([+-]?\d*\.?\d+)', line)
        y_match = re.search(r'Y=([+-]?\d*\.?\d+)', line)
        
        if x_match and y_match:
            return {
                'x': float(x_match.group(1)),
                'y': float(y_match.group(1)),
                'net': 'EXTRACTED',
                'ref': 'TP'
            }
        return None

class SchematicParser:
    def parse(self, content: str, file_path: str = "") -> Dict:
        """Parse schematic data"""
        file_ext = os.path.splitext(file_path)[1].lower() if file_path else ""
        
        if file_ext == '.schdoc':
            return self._parse_altium_schematic(content, file_path)
        elif content.strip().startswith('<?xml'):
            return self._parse_xml_schematic(content)
        else:
            return self._parse_text_netlist(content)

    def _parse_altium_schematic(self, content: str, file_path: str) -> Dict:
        """Parse Altium schematic (placeholder)"""
        return {
            'type': 'ALTIUM_SCHEMATIC',
            'file': os.path.basename(file_path),
            'nets': {},
            'components': {},
            'status': 'DETECTED_BUT_NOT_PARSED'
        }

    def _parse_xml_schematic(self, content: str) -> Dict:
        """Parse XML schematic format"""
        try:
            root = ET.fromstring(content)
            nets = {}
            components = {}
            
            # Parse nets
            for net in root.findall('.//net'):
                net_name = net.get('name') or net.get('id')
                if net_name:
                    nets[net_name] = {
                        'name': net_name,
                        'connections': []
                    }
            
            # Parse components
            for comp in root.findall('.//component') or root.findall('.//part'):
                comp_ref = comp.get('ref') or comp.get('id')
                if comp_ref:
                    components[comp_ref] = {
                        'ref': comp_ref,
                        'type': comp.get('type', 'COMPONENT'),
                        'value': comp.get('value', ''),
                        'footprint': comp.get('footprint', '')
                    }
            
            return {
                'nets': nets,
                'components': components,
                'connections': self._build_connection_matrix(nets, components)
            }
        except ET.ParseError:
            return self._parse_text_netlist(content)

    def _parse_text_netlist(self, content: str) -> Dict:
        """Parse text-based netlist"""
        nets = {}
        components = {}
        
        for line in content.strip().split('\n'):
            if line.startswith('NET'):
                net_info = self._parse_net_line(line)
                if net_info:
                    nets[net_info['name']] = net_info
            elif line.startswith('COMP'):
                comp_info = self._parse_component_line(line)
                if comp_info:
                    components[comp_info['ref']] = comp_info
        
        return {
            'nets': nets,
            'components': components,
            'connections': self._build_connection_matrix(nets, components)
        }

    def _parse_net_line(self, line: str) -> Optional[Dict]:
        """Parse a NET line from netlist"""
        parts = line.split()
        if len(parts) >= 2:
            return {'name': parts[1], 'connections': parts[2:] if len(parts) > 2 else []}
        return None

    def _parse_component_line(self, line: str) -> Optional[Dict]:
        """Parse a COMP line from netlist"""
        parts = line.split()
        if len(parts) >= 2:
            return {
                'ref': parts[1],
                'type': parts[2] if len(parts) > 2 else 'COMPONENT',
                'value': parts[3] if len(parts) > 3 else '',
                'footprint': parts[4] if len(parts) > 4 else ''
            }
        return None

    def _build_connection_matrix(self, nets: Dict, components: Dict) -> Dict:
        """Build connection matrix from nets and components"""
        return {
            'net_count': len(nets),
            'component_count': len(components),
            'status': 'BASIC_CONNECTIVITY_EXTRACTED'
        }

class RequirementsParser:
    def parse(self, content: str, file_path: str = "") -> Dict:
        """Parse requirements documents"""
        file_ext = os.path.splitext(file_path)[1].lower() if file_path else ""
        
        if file_ext == '.pdf' and PDF_SUPPORT:
            return self._parse_pdf_requirements(file_path)
        elif file_ext == '.docx' and DOCX_SUPPORT:
            return self._parse_docx_requirements(file_path)
        elif content.strip().startswith('{'):
            return json.loads(content)
        else:
            return self._parse_text_requirements(content)

    def _parse_pdf_requirements(self, file_path: str) -> Dict:
        """Parse PDF requirements document"""
        requirements = {'source': 'PDF', 'content': []}
        try:
            with open(file_path, 'rb') as f:
                pdf_reader = PyPDF2.PdfReader(f)
                text = ""
                for page in pdf_reader.pages:
                    text += page.extract_text()
                
                return self._parse_text_requirements(text)
        except Exception as e:
            return {'error': f"PDF parsing failed: {e}", 'source': 'PDF'}

    def _parse_docx_requirements(self, file_path: str) -> Dict:
        """Parse Word document requirements"""
        try:
            doc = Document(file_path)
            text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
            return self._parse_text_requirements(text)
        except Exception as e:
            return {'error': f"DOCX parsing failed: {e}", 'source': 'DOCX'}

    def _parse_text_requirements(self, content: str) -> Dict:
        """Parse text-based requirements"""
        requirements = {
            'functional': [],
            'performance': [],
            'environmental': [],
            'safety': [],
            'electrical': [],
            'mechanical': []
        }
        
        current_section = 'functional'
        for line in content.strip().split('\n'):
            line = line.strip()
            if not line:
                continue
            
            # Detect section headers
            line_lower = line.lower()
            if 'performance' in line_lower:
                current_section = 'performance'
            elif 'environmental' in line_lower:
                current_section = 'environmental'
            elif 'safety' in line_lower:
                current_section = 'safety'
            elif 'electrical' in line_lower:
                current_section = 'electrical'
            elif 'mechanical' in line_lower:
                current_section = 'mechanical'
            else:
                if line and not line.startswith('#'):
                    requirements[current_section].append(line)
        
        return requirements

def detect_file_type(filepath: str, content: str = "") -> str:
    """Detect file type based on extension and content"""
    ext = filepath.lower().split('.')[-1]
    filename = filepath.lower()

    # Altium file types
    if ext == 'bomdoc':
        return "ALTIUM_BOM"
    elif ext == 'schdoc':
        return "ALTIUM_SCHEMATIC"
    elif ext == 'pcbdoc':
        return "ALTIUM_PCB"
    elif ext == 'prjpcb':
        return "ALTIUM_PROJECT"
    
    # Standard file types
    elif ext in ['csv', 'tsv']:
        return "BOM"
    elif ext in ['xlsx', 'xls']:
        return "EXCEL_BOM"
    elif ext in ['ipc', 'net']:
        return "NETLIST"
    elif ext == 'pdf':
        return "PDF_DOCUMENT"
    elif ext in ['txt', 'md']:
        if any(keyword in filename for keyword in ['requirement', 'spec', 'hw_spec', 'sw_spec']):
            return "REQUIREMENTS"
        elif 'testpoint' in filename or 'assembly' in filename:
            return "TEST_POINTS"
        elif any(keyword in content.upper() for keyword in ['NET', 'COMPONENT', 'CONN']):
            return "NETLIST"
    elif ext == 'docx':
        return "REQUIREMENTS"
    
    # Content-based detection
    if content:
        if 'designator' in content.lower() or 'part' in content.lower():
            return "BOM"
        elif content.startswith('317') or content.startswith('327'):
            return "NETLIST"
        elif '<?xml' in content and 'schematic' in content.lower():
            return "SCHEMATIC"
    
    return "UNKNOWN"

def parse_inputs(files: List[Dict]) -> Dict[str, Any]:
    """Main parsing function"""
    parsed_data = {
        'bom_components': [],
        'test_points': [],
        'schematics': {},
        'requirements': {},
        'netlist': {},
        'metadata': {
            'total_files': len(files),
            'file_types': {},
            'parsing_errors': [],
            'supported_extensions': ['.bomdoc', '.schdoc', '.pcbdoc', '.csv', '.txt', '.pdf', '.docx', '.xml', '.ipc']
        }
    }
    
    parsers = {
        'BOM': BOMParser(),
        'ALTIUM_BOM': BOMParser(),
        'EXCEL_BOM': BOMParser(),
        'NETLIST': CoordinateParser(),
        'TEST_POINTS': CoordinateParser(),
        'ALTIUM_SCHEMATIC': SchematicParser(),
        'SCHEMATIC': SchematicParser(),
        'REQUIREMENTS': RequirementsParser(),
        'PDF_DOCUMENT': RequirementsParser()
    }
    
    for file in files:
        try:
            file_type = detect_file_type(file['path'], file.get('content', ''))
            parsed_data['metadata']['file_types'][file['path']] = file_type
            
            if file_type in parsers:
                parser = parsers[file_type]
                result = parser.parse(file.get('content', ''), file['path'])
                
                if file_type in ["BOM", "ALTIUM_BOM", "EXCEL_BOM"]:
                    parsed_data['bom_components'].extend(result)
                elif file_type in ["NETLIST", "TEST_POINTS"]:
                    parsed_data['test_points'].extend(result)
                elif file_type in ["SCHEMATIC", "ALTIUM_SCHEMATIC"]:
                    parsed_data['schematics'][file['path']] = result
                elif file_type in ["REQUIREMENTS", "PDF_DOCUMENT"]:
                    parsed_data['requirements'][file['path']] = result
                    
        except Exception as e:
            error_info = {
                'file': file['path'],
                'error': str(e),
                'type': 'parsing_error'
            }
            parsed_data['metadata']['parsing_errors'].append(error_info)
    
    # Link components to test points
    parsed_data = _link_components_to_testpoints(parsed_data)
    
    # Try to validate with Pydantic
    try:
        validated_data = PCBDataModel(**parsed_data)
        return validated_data.model_dump()
    except Exception as e:
        parsed_data['metadata']['validation_errors'] = [str(e)]
        return parsed_data

def _link_components_to_testpoints(data: Dict) -> Dict:
    """Link components to their corresponding test points"""
    component_map = {comp.designator: comp for comp in data['bom_components']}
    
    for tp in data['test_points']:
        if hasattr(tp, 'component_ref') and tp.component_ref in component_map:
            tp.linked_component = component_map[tp.component_ref]
    
    return data